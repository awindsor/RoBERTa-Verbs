#!/usr/bin/env python3
"""
Count verbs per CSV row and per group using spaCy verb lemmatization.

Input:
  - CSV with a text column (each row is processed independently)
  - Group CSV (columns=groups, rows=list lemmas)

Output:
  - CSV or XLSX with original columns + total_verb_count + per-group counts
  - JSON metadata sidecar with checksums, settings, and summary statistics

CLI Examples:
  python TextVerbGroupCounter.py input.csv groups.csv output.csv --text-col text
  python TextVerbGroupCounter.py input.csv groups.csv output.xlsx --text-col body --include-aux

Requirements:
  pip install spacy lemminflect
  python -m spacy download en_core_web_sm
  pip install PySide6  # for GUI mode only
  pip install openpyxl # for XLSX output
"""

from __future__ import annotations

import argparse
import ast
import csv
import hashlib
import json
import logging
import re
import shutil
import subprocess
import sys
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Set, Tuple

import spacy
from lemminflect import getLemma
from tqdm import tqdm

# Import openpyxl at module level to avoid threading issues
try:
    from openpyxl import Workbook
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

logger = logging.getLogger(__name__)
logger.addHandler(logging.NullHandler())


# ------------------------- version / metadata helpers -------------------------

def get_version_info() -> Dict[str, str]:
    """Return version info for this tool's dependencies."""
    try:
        import lemminflect  # type: ignore
        lemminflect_version = lemminflect.__version__
    except Exception:
        lemminflect_version = "not-installed"
    try:
        import openpyxl  # type: ignore
        openpyxl_version = openpyxl.__version__
    except Exception:
        openpyxl_version = "not-installed"

    return {
        "python": f"{sys.version_info.major}.{sys.version_info.minor}.{sys.version_info.micro}",
        "spacy": spacy.__version__,
        "lemminflect": lemminflect_version,
        "openpyxl": openpyxl_version,
    }


def compute_file_md5(path: Path) -> str:
    """Compute MD5 checksum of a file."""
    md5 = hashlib.md5()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(8192), b""):
            md5.update(chunk)
    return md5.hexdigest()


def reconstruct_command(input_csv: str, group_csv: str, output_file: str, args: argparse.Namespace) -> str:
    """Reconstruct CLI command excluding defaults to aid reproducibility."""
    cmd = ["python", "TextVerbGroupCounter.py", input_csv, group_csv, output_file]

    if args.text_col != "text":
        cmd.extend(["--text-col", args.text_col])
    if args.encoding != "utf-8":
        cmd.extend(["--encoding", args.encoding])
    if args.model != "en_core_web_sm":
        cmd.extend(["--model", args.model])
    if args.include_aux:
        cmd.append("--include-aux")
    if args.batch_size != 32:
        cmd.extend(["--batch-size", str(args.batch_size)])
    if args.where:
        cmd.extend(["--where", args.where])
    if args.load_metadata:
        cmd.extend(["--load-metadata", args.load_metadata])

    return " ".join(cmd)


def save_metadata(
    output_path: Path,
    input_path: Path,
    group_path: Path,
    input_checksum: str,
    group_checksum: str,
    output_checksum: str,
    args: Dict[str, Any],
    stats: Dict[str, Any],
    group_labels: List[str],
    command: Optional[str] = None,
) -> None:
    """Write metadata JSON next to the output file."""
    json_path = output_path.with_suffix(".json")
    metadata: Dict[str, Any] = {
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "tool": "TextVerbGroupCounter",
        "versions": get_version_info(),
        "input_file": str(input_path),
        "input_checksum": input_checksum,
        "group_file": str(group_path),
        "group_checksum": group_checksum,
        "output_file": str(output_path),
        "output_checksum": output_checksum,
        "command": command,
        "settings": {
            "text_col": args.get("text_col", "text"),
            "encoding": args.get("encoding", "utf-8"),
            "model": args.get("model", "en_core_web_sm"),
            "include_aux": args.get("include_aux", False),
            "batch_size": args.get("batch_size", 32),
            "where": args.get("where"),
            "output_format": "xlsx" if str(output_path).lower().endswith(".xlsx") else "csv",
            "group_labels": group_labels,
        },
        "statistics": stats,
    }

    with json_path.open("w", encoding="utf-8") as f:
        json.dump(metadata, f, indent=2)


def load_metadata(json_path: Path) -> Dict[str, Any]:
    """Load metadata JSON."""
    with json_path.open("r", encoding="utf-8") as f:
        data: Dict[str, Any] = json.load(f)
        return data


# ------------------------- spaCy helpers -------------------------

def load_spacy_model(model_name: str, logger: Optional[logging.Logger] = None):
    """
    Load a spaCy model, downloading it if necessary.
    """
    try:
        if logger:
            logger.info(f"Loading spaCy model: {model_name}")
        return spacy.load(model_name)
    except OSError as e:
        if "Can't find model" in str(e) or "No such file or directory" in str(e):
            if logger:
                logger.info(f"Model '{model_name}' not found locally. Downloading...")

            py_exec = sys.executable or shutil.which("python3") or shutil.which("python") or "python3"
            try:
                result = subprocess.run(
                    [py_exec, "-m", "spacy", "download", model_name],
                    capture_output=True,
                    text=True,
                    timeout=300,
                )
                if result.returncode != 0:
                    raise OSError(f"Failed to download model '{model_name}': {result.stderr}")
                if logger:
                    logger.info(f"Downloaded model '{model_name}'")
                return spacy.load(model_name)
            except subprocess.TimeoutExpired:
                raise OSError(f"Timeout downloading model '{model_name}'")
        raise


def normalize_pred_token(tok: str) -> str:
    """Normalize a token/lemma to a lemma key used by groups."""
    if tok is None:
        return ""
    s = tok.strip()
    if not s:
        return ""
    if s.startswith("Ġ"):
        s = s[1:]
    if s.startswith("▁"):
        s = s[1:]
    s = s.strip().lower()
    if not s:
        return ""
    try:
        lemmas = getLemma(s, "VERB")
        if lemmas:
            return (lemmas[0] or s).lower()
    except Exception:
        pass
    return s


def load_groups_csv(group_csv_path: str, encoding: str, logger: logging.Logger) -> Tuple[List[str], Dict[str, Set[str]]]:
    """Load groups from a wide CSV (columns=groups, rows=list lemmas) into mapping."""
    p = Path(group_csv_path)
    if not p.exists():
        raise FileNotFoundError(f"Group CSV not found: {p}")
    with p.open("r", newline="", encoding=encoding) as f:
        reader = csv.reader(f)
        header = next(reader, None)
        if not header:
            raise ValueError("Group CSV is empty or missing header row")
        group_labels: List[str] = [h.strip() for h in header if h and h.strip()]
        lemma_to_groups: Dict[str, Set[str]] = {}
        for row in reader:
            for idx, g in enumerate(group_labels):
                if idx >= len(row):
                    continue
                cell = (row[idx] or "").strip()
                if not cell:
                    continue
                key = normalize_pred_token(cell)
                if not key:
                    continue
                s = lemma_to_groups.get(key)
                if s is None:
                    s = set()
                    lemma_to_groups[key] = s
                s.add(g)
    logger.info(f"Loaded {len(group_labels)} groups; {len(lemma_to_groups)} unique lemmas mapped")
    return group_labels, lemma_to_groups


# ------------------------- output helpers -------------------------

def _bool_from_row_expression(expr: str, row: Dict[str, str]) -> bool:
    """
    Evaluate a boolean expression after replacing {{col}} placeholders.
    Example: {{status}} == "ok" and {{score}} == "10"
    """
    def replace_placeholder(match: re.Match) -> str:
        col = match.group(1)
        val = row.get(col, "")
        return repr(val)

    replaced = re.sub(r"\{\{\s*([^}]+?)\s*\}\}", replace_placeholder, expr)
    try:
        parsed = ast.parse(replaced, mode="eval")
        return bool(_safe_eval_bool(parsed))
    except Exception as e:
        raise ValueError(f"Invalid --where expression: {expr!r} -> {replaced!r} ({e})")


def _safe_eval_bool(node: ast.AST) -> Any:
    """Safely evaluate a restricted boolean expression AST."""
    if isinstance(node, ast.Expression):
        return _safe_eval_bool(node.body)
    if isinstance(node, ast.Constant):
        return node.value
    if isinstance(node, ast.Tuple):
        return tuple(_safe_eval_bool(elt) for elt in node.elts)
    if isinstance(node, ast.List):
        return [_safe_eval_bool(elt) for elt in node.elts]
    if isinstance(node, ast.Set):
        return {_safe_eval_bool(elt) for elt in node.elts}
    if isinstance(node, ast.Dict):
        return {_safe_eval_bool(k): _safe_eval_bool(v) for k, v in zip(node.keys or [], node.values or []) if k is not None}
    if isinstance(node, ast.UnaryOp) and isinstance(node.op, ast.Not):
        return not _safe_eval_bool(node.operand)
    if isinstance(node, ast.BoolOp):
        if isinstance(node.op, ast.And):
            return all(_safe_eval_bool(v) for v in node.values)
        if isinstance(node.op, ast.Or):
            return any(_safe_eval_bool(v) for v in node.values)
    if isinstance(node, ast.Compare):
        left = _safe_eval_bool(node.left)
        for op, comparator in zip(node.ops, node.comparators):
            right = _safe_eval_bool(comparator)
            if isinstance(op, ast.Eq):
                ok = left == right
            elif isinstance(op, ast.NotEq):
                ok = left != right
            elif isinstance(op, ast.Lt):
                ok = left < right
            elif isinstance(op, ast.LtE):
                ok = left <= right
            elif isinstance(op, ast.Gt):
                ok = left > right
            elif isinstance(op, ast.GtE):
                ok = left >= right
            elif isinstance(op, ast.In):
                ok = left in right
            elif isinstance(op, ast.NotIn):
                ok = left not in right
            else:
                raise ValueError(f"Unsupported comparison operator: {type(op).__name__}")
            if not ok:
                return False
            left = right
        return True
    raise ValueError(f"Unsupported expression element: {type(node).__name__}")


def _ensure_unique_col(name: str, existing: Iterable[str]) -> str:
    existing_set = set(existing)
    if name not in existing_set:
        return name
    suffix = 1
    while True:
        candidate = f"{name}_{suffix}"
        if candidate not in existing_set:
            return candidate
        suffix += 1


def build_group_count_columns(fieldnames: Sequence[str], group_labels: List[str]) -> Tuple[List[str], Dict[str, str]]:
    """Create safe output column names for each group and a mapping from group->column."""
    existing = set(fieldnames)
    out_cols: List[str] = []
    col_map: Dict[str, str] = {}
    for g in group_labels:
        base = f"group_count_{g}"
        name = base
        suffix = 1
        while name in existing or name in out_cols:
            name = f"{base}_{suffix}"
            suffix += 1
        out_cols.append(name)
        col_map[g] = name
    return out_cols, col_map


def write_xlsx(output_path: Path, header: List[str], rows: Iterable[List[Any]]) -> None:
    if not OPENPYXL_AVAILABLE:
        raise SystemExit(
            "Excel output requested but openpyxl is not installed. Install with: pip install openpyxl"
        )
    wb = Workbook()
    ws = wb.active
    ws.title = "verb_group_counts"
    ws.append(header)
    for row in rows:
        ws.append(row)
    wb.save(str(output_path))


# ------------------------- core processing -------------------------

def process_rows(
    input_path: Path,
    group_path: Path,
    output_path: Path,
    *,
    text_col: str,
    encoding: str,
    model: str,
    include_aux: bool,
    batch_size: int,
    where: Optional[str],
    on_progress=None,
    on_progress_value=None,
    total_rows_expected: Optional[int] = None,
) -> Dict[str, Any]:
    """Process input CSV rows, write output, and return stats."""
    start = time.time()
    group_labels, lemma_to_groups = load_groups_csv(str(group_path), encoding, logger)

    nlp = load_spacy_model(model, logger)
    components_to_enable = [
        c
        for c in ["transformer", "tok2vec", "tagger", "parser", "senter", "attribute_ruler", "lemmatizer"]
        if c in nlp.pipe_names
    ]
    if components_to_enable:
        nlp.select_pipes(enable=components_to_enable)
    if "parser" not in nlp.pipe_names:
        if "sentencizer" not in nlp.pipe_names:
            nlp.add_pipe("sentencizer")
            logger.info("Parser not found in pipeline; added sentencizer for sentence boundaries.")

    rows_processed = 0
    rows_filtered_out = 0
    total_verbs = 0
    total_group_matches = 0
    rows_with_verbs = 0
    rows_with_group_hits = 0

    def emit(msg: str):
        if on_progress:
            on_progress(msg)

    emit(f"Reading input CSV: {input_path}")
    with input_path.open("r", newline="", encoding=encoding, errors="replace") as f:
        reader = csv.DictReader(f)
        if not reader.fieldnames:
            raise ValueError("Input CSV is missing a header row")
        if text_col not in reader.fieldnames:
            raise ValueError(f"Missing text column {text_col!r} in input CSV")

        fieldnames = list(reader.fieldnames)
        total_col = _ensure_unique_col("total_verb_count", fieldnames)
        group_cols, group_col_map = build_group_count_columns(fieldnames + [total_col], group_labels)
        out_fieldnames = fieldnames + [total_col] + group_cols

        xlsx_buffer: Optional[List[Dict[str, Any]]] = None
        if output_path.suffix.lower() == ".csv":
            out_file = output_path.open("w", newline="", encoding=encoding)
            writer = csv.DictWriter(out_file, fieldnames=out_fieldnames)
            writer.writeheader()
            close_out = out_file.close
            xlsx_buffer = None
        elif output_path.suffix.lower() == ".xlsx":
            out_file = None
            writer = None
            def close_out() -> None:
                return None
            xlsx_buffer = []
        else:
            raise SystemExit("Output must end with .csv or .xlsx")

        try:
            emit("Processing rows...")
            batch: List[Tuple[str, Dict[str, str]]] = []
            pbar = tqdm(unit=" rows", desc="Processing", leave=True)
            for row in reader:
                if where:
                    if not _bool_from_row_expression(where, row):
                        rows_filtered_out += 1
                        continue
                text = row.get(text_col, "") or ""
                batch.append((text, row))
                if len(batch) >= batch_size:
                    (
                        rows_processed,
                        total_verbs,
                        total_group_matches,
                        rows_with_verbs,
                        rows_with_group_hits,
                    ) = _process_batch(
                        nlp,
                        batch,
                        group_labels,
                        lemma_to_groups,
                        group_col_map,
                        total_col,
                        include_aux,
                        writer,
                        xlsx_buffer,
                        rows_processed,
                        total_verbs,
                        total_group_matches,
                        rows_with_verbs,
                        rows_with_group_hits,
                        total_rows_expected,
                        on_progress_value,
                    )
                    pbar.update(len(batch))
                    batch = []

            if batch:
                (
                    rows_processed,
                    total_verbs,
                    total_group_matches,
                    rows_with_verbs,
                    rows_with_group_hits,
                ) = _process_batch(
                    nlp,
                    batch,
                    group_labels,
                    lemma_to_groups,
                    group_col_map,
                    total_col,
                    include_aux,
                    writer,
                    xlsx_buffer,
                    rows_processed,
                    total_verbs,
                    total_group_matches,
                    rows_with_verbs,
                    rows_with_group_hits,
                    total_rows_expected,
                    on_progress_value,
                )
                pbar.update(len(batch))
            pbar.close()

            if output_path.suffix.lower() == ".xlsx":
                if xlsx_buffer is None:
                    raise RuntimeError("xlsx_buffer is missing for XLSX output")
                emit(f"Writing XLSX file: {output_path}")
                rows_as_lists = [
                    [row.get(h, "") for h in out_fieldnames] for row in xlsx_buffer
                ]
                write_xlsx(output_path, out_fieldnames, rows_as_lists)
        finally:
            close_out()

    elapsed = time.time() - start
    stats = {
        "total_rows": rows_processed,
        "rows_filtered_out": rows_filtered_out,
        "rows_with_verbs": rows_with_verbs,
        "rows_with_group_hits": rows_with_group_hits,
        "total_verbs": total_verbs,
        "total_group_matches": total_group_matches,
        "elapsed_seconds": elapsed,
    }
    emit("\n=== Processing Complete ===")
    emit(f"Total rows processed: {rows_processed}")
    emit(f"Rows filtered out: {rows_filtered_out}")
    emit(f"Rows with verbs: {rows_with_verbs}")
    emit(f"Rows with group matches: {rows_with_group_hits}")
    emit(f"Total verbs found: {total_verbs}")
    emit(f"Total group matches: {total_group_matches}")
    emit(f"Time elapsed: {elapsed:.2f}s")
    emit(f"Output written to: {output_path}")
    return stats


def _process_batch(
    nlp,
    batch: List[Tuple[str, Dict[str, str]]],
    group_labels: List[str],
    lemma_to_groups: Dict[str, Set[str]],
    group_col_map: Dict[str, str],
    total_col: str,
    include_aux: bool,
    writer,
    xlsx_buffer,
    rows_processed: int,
    total_verbs: int,
    total_group_matches: int,
    rows_with_verbs: int,
    rows_with_group_hits: int,
    total_rows_expected: Optional[int],
    on_progress_value=None,
) -> Tuple[int, int, int, int, int]:
    for doc, row in nlp.pipe(batch, as_tuples=True):
        row_total_verbs = 0
        group_counts = {g: 0 for g in group_labels}

        for tok in doc:
            is_verb = tok.pos_ == "VERB" or (include_aux and tok.pos_ == "AUX")
            if not is_verb:
                continue
            row_total_verbs += 1
            lemma = normalize_pred_token(tok.lemma_)
            if not lemma:
                continue
            groups = lemma_to_groups.get(lemma)
            if not groups:
                continue
            for g in groups:
                group_counts[g] += 1

        out_row = dict(row)
        out_row[total_col] = row_total_verbs
        for g in group_labels:
            out_row[group_col_map[g]] = group_counts[g]

        if writer is not None:
            writer.writerow(out_row)
        else:
            xlsx_buffer.append(out_row)

        rows_processed += 1
        total_verbs += out_row[total_col]
        row_group_hits = 0
        for g in group_labels:
            row_group_hits += out_row[group_col_map[g]]
        total_group_matches += row_group_hits
        if out_row[total_col] > 0:
            rows_with_verbs += 1
        if row_group_hits > 0:
            rows_with_group_hits += 1
        if on_progress_value:
            on_progress_value(rows_processed, total_rows_expected or 0)
    return rows_processed, total_verbs, total_group_matches, rows_with_verbs, rows_with_group_hits


# ------------------------- CLI -------------------------

def run_cli(argv: Optional[List[str]] = None, on_progress=None, on_progress_value=None, total_rows_expected=None) -> None:
    ap = argparse.ArgumentParser(description="Count verbs per row and per group from a CSV text column")
    ap.add_argument("input_csv", help="Input CSV with a text column")
    ap.add_argument("group_csv", help="Group CSV (columns=groups, rows=list lemmas)")
    ap.add_argument("output", help="Output file (.csv or .xlsx)")
    ap.add_argument("--text-col", default="text", help="Text column name (default: text)")
    ap.add_argument("--encoding", default="utf-8", help="File encoding (default: utf-8)")
    ap.add_argument("--model", default="en_core_web_sm", help="spaCy model (default: en_core_web_sm)")
    ap.add_argument("--include-aux", action="store_true", help="Count AUX tokens as verbs")
    ap.add_argument("--batch-size", type=int, default=32, help="spaCy pipe batch size (default: 32)")
    ap.add_argument("--where", help="Row filter expression using {{column}} placeholders")
    ap.add_argument("--load-metadata", help="Load settings from a previous JSON metadata file")

    args = ap.parse_args(argv)

    if args.load_metadata:
        meta = load_metadata(Path(args.load_metadata))
        settings = meta.get("settings", {}) if isinstance(meta.get("settings"), dict) else {}
        args.text_col = settings.get("text_col", args.text_col)
        args.encoding = settings.get("encoding", args.encoding)
        args.model = settings.get("model", args.model)
        args.include_aux = settings.get("include_aux", args.include_aux)
        args.batch_size = settings.get("batch_size", args.batch_size)
        args.where = settings.get("where", args.where)

    input_path = Path(args.input_csv)
    group_path = Path(args.group_csv)
    output_path = Path(args.output)

    if not input_path.exists():
        raise SystemExit(f"Input CSV not found: {input_path}")
    if not group_path.exists():
        raise SystemExit(f"Group CSV not found: {group_path}")

    if output_path.suffix.lower() not in {".csv", ".xlsx"}:
        raise SystemExit("Output must end with .csv or .xlsx")

    stats = process_rows(
        input_path,
        group_path,
        output_path,
        text_col=args.text_col,
        encoding=args.encoding,
        model=args.model,
        include_aux=args.include_aux,
        batch_size=args.batch_size,
        where=args.where,
        on_progress=on_progress,
        on_progress_value=on_progress_value,
        total_rows_expected=total_rows_expected,
    )

    input_checksum = compute_file_md5(input_path)
    group_checksum = compute_file_md5(group_path)
    output_checksum = compute_file_md5(output_path)

    command = reconstruct_command(str(input_path), str(group_path), str(output_path), args)
    save_metadata(
        output_path=output_path,
        input_path=input_path,
        group_path=group_path,
        input_checksum=input_checksum,
        group_checksum=group_checksum,
        output_checksum=output_checksum,
        args={
            "text_col": args.text_col,
            "encoding": args.encoding,
            "model": args.model,
            "include_aux": args.include_aux,
            "batch_size": args.batch_size,
            "where": args.where,
        },
        stats=stats,
        group_labels=load_groups_csv(str(group_path), args.encoding, logger)[0],
        command=command,
    )


# ------------------------- GUI -------------------------

def run_gui() -> None:
    try:
        from PySide6.QtCore import QThread, Signal
        from PySide6.QtWidgets import (
            QApplication,
            QFileDialog,
            QGridLayout,
            QGroupBox,
            QHBoxLayout,
            QLabel,
            QLineEdit,
            QMainWindow,
            QMessageBox,
            QProgressBar,
            QPushButton,
            QSpinBox,
            QTextEdit,
            QVBoxLayout,
            QWidget,
            QCheckBox,
        )
    except ImportError:
        print("Error: PySide6 is required for GUI mode. Install with: pip install PySide6")
        sys.exit(1)

    class Worker(QThread):
        progress = Signal(str)
        progress_value = Signal(int, int)
        finished = Signal(bool, str)

        def __init__(
            self,
            input_path: Path,
            group_path: Path,
            output_path: Path,
            text_col: str,
            encoding: str,
            model: str,
            include_aux: bool,
            batch_size: int,
            where: Optional[str],
            metadata_path: Optional[Path],
        ):
            super().__init__()
            self.input_path = input_path
            self.group_path = group_path
            self.output_path = output_path
            self.text_col = text_col
            self.encoding = encoding
            self.model = model
            self.include_aux = include_aux
            self.batch_size = batch_size
            self.where = where
            self.metadata_path = metadata_path

        def run(self):
            try:
                self.progress.emit("Starting processing...")
                total_rows = 0
                try:
                    with self.input_path.open(newline="", encoding=self.encoding) as fin:
                        reader = csv.reader(fin)
                        next(reader, None)
                        for _ in reader:
                            total_rows += 1
                    self.progress_value.emit(0, total_rows)
                except Exception:
                    total_rows = 0
                    self.progress_value.emit(0, 0)

                argv = [
                    str(self.input_path),
                    str(self.group_path),
                    str(self.output_path),
                    "--text-col",
                    self.text_col,
                    "--encoding",
                    self.encoding,
                    "--model",
                    self.model,
                    "--batch-size",
                    str(self.batch_size),
                ]
                if self.include_aux:
                    argv.append("--include-aux")
                if self.where:
                    argv.extend(["--where", self.where])
                if self.metadata_path:
                    argv.extend(["--load-metadata", str(self.metadata_path)])

                self.progress.emit(f"Input: {argv[0]}")
                self.progress.emit(f"Group: {argv[1]}")
                self.progress.emit(f"Output: {argv[2]}")

                def progress_callback(msg):
                    if msg is not None:
                        self.progress.emit(msg)

                def progress_value_callback(current, total):
                    self.progress_value.emit(current, total if total is not None else total_rows)

                run_cli(
                    argv,
                    on_progress=progress_callback,
                    on_progress_value=progress_value_callback,
                    total_rows_expected=total_rows if total_rows > 0 else None,
                )
                self.progress.emit("✓ Processing complete")
                self.finished.emit(True, "Completed successfully")
            except SystemExit as e:
                msg = f"Exit code {e.code}: Check the log for details"
                self.progress.emit(f"✗ {msg}")
                self.finished.emit(False, msg)
            except Exception as e:
                import traceback
                msg = f"Error: {str(e)}"
                self.progress.emit(f"✗ {msg}")
                self.progress.emit(traceback.format_exc())
                self.finished.emit(False, msg)

    class MainWindow(QMainWindow):
        def __init__(self):
            super().__init__()
            self.setWindowTitle("Text → Verb Group Counts")

            self.input_edit = QLineEdit()
            self.group_edit = QLineEdit()
            self.output_edit = QLineEdit()
            self.text_col_edit = QLineEdit("text")
            self.encoding_edit = QLineEdit("utf-8")
            self.model_edit = QLineEdit("en_core_web_sm")
            self.where_edit = QLineEdit()
            self.include_aux_check = QCheckBox("Include AUX tokens as verbs")
            self.batch_spin = QSpinBox()
            self.batch_spin.setRange(1, 512)
            self.batch_spin.setValue(32)
            self.meta_path: Optional[Path] = None

            self.progress_bar = QProgressBar()
            self.progress_bar.setRange(0, 100)
            self.progress_bar.setValue(0)
            self.progress_bar.setTextVisible(True)
            self.progress_bar.setFormat("Idle")
            self.progress_label = QLabel("Ready")
            self.progress_label.setStyleSheet("color: #666; font-style: italic;")
            self.progress_text = QTextEdit()
            self.progress_text.setReadOnly(True)
            self.progress_text.setMaximumHeight(200)

            browse_in = QPushButton("Browse...")
            browse_groups = QPushButton("Browse...")
            browse_out = QPushButton("Browse...")
            load_json_btn = QPushButton("Load Settings from JSON")
            run_btn = QPushButton("Run")

            browse_in.clicked.connect(self.pick_input)
            browse_groups.clicked.connect(self.pick_groups)
            browse_out.clicked.connect(self.pick_output)
            load_json_btn.clicked.connect(self.pick_metadata)
            run_btn.clicked.connect(self.start)

            file_grid = QGridLayout()
            file_grid.addWidget(QLabel("Input CSV:"), 0, 0)
            file_grid.addWidget(self.input_edit, 0, 1)
            file_grid.addWidget(browse_in, 0, 2)
            file_grid.addWidget(QLabel("Group CSV:"), 1, 0)
            file_grid.addWidget(self.group_edit, 1, 1)
            file_grid.addWidget(browse_groups, 1, 2)
            file_grid.addWidget(QLabel("Output File:"), 2, 0)
            file_grid.addWidget(self.output_edit, 2, 1)
            file_grid.addWidget(browse_out, 2, 2)

            settings_box = QGroupBox("Settings")
            settings_layout = QVBoxLayout()

            metadata_layout = QHBoxLayout()
            metadata_layout.addWidget(load_json_btn)
            metadata_layout.addStretch()
            settings_layout.addLayout(metadata_layout)

            options_layout = QHBoxLayout()

            text_col_layout = QVBoxLayout()
            text_col_layout.addWidget(QLabel("Text Column:"))
            text_col_layout.addWidget(self.text_col_edit)

            encoding_layout = QVBoxLayout()
            encoding_layout.addWidget(QLabel("Encoding:"))
            encoding_layout.addWidget(self.encoding_edit)

            model_layout = QVBoxLayout()
            model_layout.addWidget(QLabel("spaCy Model:"))
            model_layout.addWidget(self.model_edit)

            where_layout = QVBoxLayout()
            where_layout.addWidget(QLabel("Row Filter ({{col}} syntax):"))
            where_layout.addWidget(self.where_edit)

            batch_layout = QVBoxLayout()
            batch_layout.addWidget(QLabel("Batch Size:"))
            batch_layout.addWidget(self.batch_spin)

            options_layout.addLayout(text_col_layout)
            options_layout.addLayout(encoding_layout)
            options_layout.addLayout(model_layout)
            options_layout.addLayout(where_layout)
            options_layout.addLayout(batch_layout)

            settings_layout.addLayout(options_layout)
            settings_layout.addWidget(self.include_aux_check)
            settings_box.setLayout(settings_layout)

            main = QVBoxLayout()
            main.addLayout(file_grid)
            main.addWidget(settings_box)
            main.addWidget(run_btn)
            main.addWidget(QLabel("Progress"))
            main.addWidget(self.progress_bar)
            main.addWidget(self.progress_label)
            main.addWidget(self.progress_text)

            container = QWidget()
            container.setLayout(main)
            self.setCentralWidget(container)

            self.worker: Optional[Worker] = None

        def pick_input(self):
            path, _ = QFileDialog.getOpenFileName(self, "Select input CSV", "", "CSV Files (*.csv)")
            if path:
                self.input_edit.setText(path)

        def pick_groups(self):
            path, _ = QFileDialog.getOpenFileName(self, "Select group CSV", "", "CSV Files (*.csv)")
            if path:
                self.group_edit.setText(path)

        def pick_output(self):
            path, _ = QFileDialog.getSaveFileName(self, "Select output", "", "CSV/XLSX Files (*.csv *.xlsx)")
            if path:
                self.output_edit.setText(path)

        def pick_metadata(self):
            path, _ = QFileDialog.getOpenFileName(self, "Select metadata JSON", "", "JSON Files (*.json)")
            if not path:
                return
            self.meta_path = Path(path)
            try:
                meta = load_metadata(Path(path))
                tool = meta.get("tool")
                if tool != "TextVerbGroupCounter":
                    QMessageBox.warning(self, "Metadata error", f"Unknown tool in metadata: {tool}")
                    return
                input_file = meta.get("input_file")
                if input_file:
                    self.input_edit.setText(str(input_file))
                group_file = meta.get("group_file")
                if group_file:
                    self.group_edit.setText(str(group_file))
                output_file = meta.get("output_file")
                if output_file:
                    self.output_edit.setText(str(output_file))
                settings = meta.get("settings", {}) if isinstance(meta.get("settings"), dict) else {}
                text_col = settings.get("text_col")
                if text_col:
                    self.text_col_edit.setText(str(text_col))
                encoding = settings.get("encoding")
                if encoding:
                    self.encoding_edit.setText(str(encoding))
                model = settings.get("model")
                if model:
                    self.model_edit.setText(str(model))
                include_aux = settings.get("include_aux")
                self.include_aux_check.setChecked(bool(include_aux))
                batch_size = settings.get("batch_size")
                if isinstance(batch_size, int) and batch_size > 0:
                    self.batch_spin.setValue(batch_size)
                where = settings.get("where")
                if where:
                    self.where_edit.setText(str(where))
                QMessageBox.information(self, "Metadata loaded", "✓ Loaded TextVerbGroupCounter metadata")
            except Exception as e:
                QMessageBox.warning(self, "Metadata error", f"Could not read metadata JSON: {str(e)}")

        def start(self):
            input_path = Path(self.input_edit.text().strip())
            group_path = Path(self.group_edit.text().strip())
            output_path = Path(self.output_edit.text().strip())
            if not input_path.exists():
                QMessageBox.warning(self, "Input missing", "Select a valid input CSV")
                return
            if not group_path.exists():
                QMessageBox.warning(self, "Group missing", "Select a valid group CSV")
                return
            if input_path.suffix.lower() != ".csv":
                QMessageBox.warning(self, "Invalid input", "Input file must be a .csv file")
                return
            if not self.output_edit.text().strip():
                QMessageBox.warning(self, "Output missing", "Select an output filename")
                return
            if output_path.suffix.lower() not in {".csv", ".xlsx"}:
                QMessageBox.warning(self, "Bad output", "Output must end with .csv or .xlsx")
                return

            text_col = self.text_col_edit.text().strip() or "text"
            encoding = self.encoding_edit.text().strip() or "utf-8"
            model = self.model_edit.text().strip() or "en_core_web_sm"
            include_aux = self.include_aux_check.isChecked()
            batch_size = int(self.batch_spin.value())
            where = self.where_edit.text().strip() or None

            self.progress_text.clear()
            self.progress_text.append("Starting worker thread...")
            self.progress_bar.setRange(0, 1)
            self.progress_bar.setValue(0)
            self.progress_bar.setFormat("Starting...")
            self.worker = Worker(
                input_path=input_path,
                group_path=group_path,
                output_path=output_path,
                text_col=text_col,
                encoding=encoding,
                model=model,
                include_aux=include_aux,
                batch_size=batch_size,
                where=where,
                metadata_path=self.meta_path,
            )
            self.worker.progress.connect(self.append_log)
            self.worker.progress_value.connect(self.update_progress)
            self.worker.finished.connect(self.done)
            self.worker.start()

        def append_log(self, msg: str):
            self.progress_text.append(msg)

        def update_progress(self, current: int, total: int):
            if total and total > 0:
                if self.progress_bar.maximum() != total:
                    self.progress_bar.setRange(0, total)
                self.progress_bar.setValue(current)
                self.progress_bar.setFormat(f"{current}/{total}")
            else:
                if self.progress_bar.maximum() != 1:
                    self.progress_bar.setRange(0, 1)
                self.progress_bar.setValue(0)
                self.progress_bar.setFormat("Working...")

        def done(self, ok: bool, msg: str):
            self.progress_bar.setRange(0, 100)
            self.progress_bar.setValue(100)
            self.progress_bar.setFormat("Done" if ok else "Error")
            self.progress_label.setText("Complete" if ok else "Failed")
            if ok:
                QMessageBox.information(self, "Done", msg)
            else:
                QMessageBox.critical(self, "Error", msg)

    app = QApplication(sys.argv)
    window = MainWindow()
    window.resize(720, 520)
    window.show()
    app.exec()


def main() -> None:
    if len(sys.argv) == 1:
        run_gui()
    else:
        run_cli()


if __name__ == "__main__":
    main()
